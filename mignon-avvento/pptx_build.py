"""Ricostruisce le tavole come PPTX nativi: caselle di testo e forme vere,
non immagini. Canva, PowerPoint, Keynote e Google Slides li importano
mantenendo tutto modificabile. Un file per formato, perche un PPTX ha una
sola dimensione di pagina.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from copy import deepcopy
import lxml.etree as etree

PX = 9525  # 1 px @96dpi in EMU


def C(h):
    return RGBColor.from_string(h)

INK, INK2, INK3 = "14100E", "1C1613", "1E1714"
PAPER, CREAM, WHITE = "F3ECE2", "F7F1E6", "FFFFFF"
RED, ORANGE, PURPLE, YELLOW = "E42313", "F08019", "7A2E8E", "F6D01A"
MUTED_D, DIM_D, SOFT_D = "9A8F82", "B0A498", "C9BEB1"
MUTED_L, DIM_L = "6B6157", "5A5148"

SERIF, SANS = "Bodoni Moda", "Jost"
BLACK, GROTESK = "Archivo Black", "Archivo"
SERIF_C, SANS_M = "Cormorant Garamond", "Mulish"


def deck(w, h):
    p = Presentation()
    p.slide_width, p.slide_height = Emu(w * PX), Emu(h * PX)
    return p


def page(prs, bg):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = C(bg); r.line.fill.background()
    r.shadow.inherit = False
    return s


def box(s, x, y, w, h, fill=None, line=None, lw=1):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = C(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = C(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def txt(s, x, y, w, h, parts, size, font=SANS, color="000000", align="l",
        bold=False, italic=False, spacing=None, leading=1.25, anchor="t"):
    """parts: stringa, oppure lista di (testo, {override}) per run misti."""
    tb = s.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    lines = parts if isinstance(parts, list) else [parts]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.line_spacing = leading
        runs = line if isinstance(line, list) else [(line, {})]
        for text, ov in runs:
            r = p.add_run(); r.text = text
            f = r.font
            f.name = ov.get("font", font)
            f.size = Pt(ov.get("size", size) * 0.75)  # px -> pt
            f.bold = ov.get("bold", bold)
            f.italic = ov.get("italic", italic)
            f.color.rgb = C(ov.get("color", color))
            sp = ov.get("spacing", spacing)  # in px
            if sp:
                f._rPr.set("spc", str(int(sp * 0.75 * 100)))
    return tb


def pic(s, path, x, y, h):
    from PIL import Image
    iw, ih = Image.open(path).size
    return s.shapes.add_picture(path, Emu(x * PX), Emu(y * PX), height=Emu(h * PX),
                                width=Emu(round(h * iw / ih) * PX))


def rule(s, x, y, w, color, thick=1):
    return box(s, x, y, w, thick, fill=color)
